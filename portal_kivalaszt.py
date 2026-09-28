import re
from pathlib import Path
import openpyxl
import pandas as pd

_HONAP_ROVIDITES = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "maj": 5, "jun": 6,
    "jul": 7, "aug": 8, "sze": 9, "okt": 10, "nov": 11, "dec": 12,
}
def fajlnev_datum(fajlnev:str)->tuple[int, int]:
    illesztes=re.match(r"(\d{2})\s+([a-z]{3})", fajlnev.lower())
    if not illesztes:
        raise ValueError(f"ismeretlen fajlnév formátum: {fajlnev}")
    ev=2000+int(illesztes.group(1))
    honap=_HONAP_ROVIDITES[illesztes.group(2)]
    return ev, honap

def jelentes_info_beolvasas(munkafuzet:openpyxl.Workbook)->dict:
    lap=munkafuzet['Report info']
    info={}
    for kulcs,ertek in lap.iter_rows(min_row=1, values_only=True):
        if kulcs is None:
            continue
        info[kulcs]=ertek
    return info

def jelentes_osszesit(mappa:str|Path)->pd.DataFrame:
    mappa=Path(mappa)
    sorok=[]

    for fajl in sorted(mappa.glob("*.xlsx")):
        ev_fajlnev,honap_fajlnev=fajlnev_datum(fajl.name)
        munkafuzet=openpyxl.load_workbook(fajl,data_only=True)

        info=jelentes_info_beolvasas(munkafuzet)
        lap=munkafuzet['Report data']
        fejlec,*adat_sorok=lap.iter_rows(values_only=True)

        for csatorna, views, visits in adat_sorok:
            if csatorna is None:
                continue
            sorok.append({
                "fajlnev": fajl.name,
                "ev": ev_fajlnev,
                "honap": honap_fajlnev,
                "riport_idoszak": info.get("Time period"),
                "riport_forras": info.get("Source"),
                "riport_letrehozva": info.get("Created"),
                "csatorna": csatorna,
                "views": views,
                "visits": visits,
            })

    return pd.DataFrame(sorok)

df = jelentes_osszesit(r"nyers/weboldal látogatottság")
print(df.shape)
print(df.head(20).to_string())
print(df["fajlnev"].nunique(), "fájl,", df["csatorna"].nunique(), "csatorna")
kimenet_osszesitve = Path(r"nyers/weboldal látogatottság/latogatottsag_osszesitve.csv")
df.to_csv(kimenet_osszesitve, index=False, encoding="iso-8859-2")
print("Mentve:", kimenet_osszesitve)


def top_20_nezettseg(osszesitett_csv: str | Path, n: int = 20) -> pd.DataFrame:
    df = pd.read_csv(osszesitett_csv, encoding="iso-8859-2")
    df = df[df["csatorna"] != "All"].copy()
    df["views"] = pd.to_numeric(df["views"], errors="coerce")
    rangsor = (
        df.groupby("csatorna", as_index=False)["views"]
        .sum()
        .sort_values("views", ascending=False)
        .head(n)
        .reset_index(drop=True)
    )
    rangsor.index += 1
    return rangsor

top20 = top_20_nezettseg(kimenet_osszesitve, n=20)
print(top20.to_string())
kimenet_top20 = Path(r"nyers/weboldal látogatottság/latogatottsag_top20.csv")
top20.to_csv(kimenet_top20, encoding="iso-8859-2")
print("Mentve:", kimenet_top20)